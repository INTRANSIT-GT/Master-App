#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Carga las plantillas de producción (Plantillas_Master.xlsx) en el esquema master.

  python cargar_plantillas.py                      → simula: inserta todo y REVIERTE (valida de verdad)
  python cargar_plantillas.py --aplicar            → inserta y confirma (todo o nada)
  python cargar_plantillas.py --probar             → usa la fila de EJEMPLO de cada hoja, revierte y resume
  python cargar_plantillas.py --tabla viajes --simular
  python cargar_plantillas.py --reemplazar --simular          → borra los datos de negocio, carga el libro y revierte
  python cargar_plantillas.py --reemplazar --aplicar --si-reemplazar
  --archivo <ruta.xlsx>   --verbose

Reglas: si una fila falla se revierte solo esa fila (savepoint) y se sigue; al final, con --aplicar
solo se confirma si no hubo ningún error. Antes de confirmar recalcula master.secuencias.
"""
import argparse, json, os, re, shutil, sys, uuid
from datetime import date, datetime, time as dtime
from decimal import Decimal, InvalidOperation
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import psycopg
from psycopg.types.json import Jsonb
from openpyxl import load_workbook

BASE = os.path.dirname(os.path.abspath(__file__))
CONEXION = os.environ.get("DATABASE_URL_FILE", os.path.join(os.path.expanduser("~"), "GensparkCode", "_datos", "conexion.txt"))
RESPALDOS = os.environ.get("RESPALDOS_DIR", os.path.join(os.path.expanduser("~"), "GensparkCode", "_datos", "respaldos"))

ORDEN = ["clientes", "clientes_fiscal", "flota", "pilotos", "rutas", "tarifas", "mem",
         "proveedores", "viajes", "viajes_tarifa", "viajes_costos", "viajes_viaticos", "liquidaciones",
         "agrupaciones", "pagos", "ordenes_pago", "perfiles", "acceso_cliente", "gastos", "mantenimientos"]
SECUENCIA = ORDEN[:ORDEN.index("proveedores") + 1] + ["catalogos"] + ORDEN[ORDEN.index("proveedores") + 1:]
LISTAS = {"catalogos": None}
CLAVES_LISTA = {"tipoUnidad", "opl", "tipoCarga", "tipoViaje", "arrastre"}
SALTAR_EN_PRUEBA = {"perfiles", "acceso_cliente"}

# Orden de borrado para empezar de cero: primero los hijos (llaves foráneas), al final los padres.
PURGA = ["evidencias", "notificaciones", "viajes_costos", "viajes_tarifa", "viajes_viaticos", "liquidaciones", "pagos",
         "ordenes_pago", "agrupaciones", "viajes", "tarifas", "rutas", "pilotos", "flota",
         "clientes_fiscal", "destinos", "acceso_cliente", "clientes", "mem", "proveedores",
         "gastos", "mantenimientos"]

PROBE_FIRST = {"clientes": "CLI-EJ", "clientes_fiscal": "CLI-EJ", "flota": "EJ-001",
               "pilotos": "P-EJ1", "rutas": "R-EJ1", "tarifas": "TAR-990001", "mem": "2027-01-01",
               "proveedores": "OPL-EJ",
               "viajes": "200099", "viajes_tarifa": "200099", "viajes_costos": "200099",
               "viajes_viaticos": "200099", "liquidaciones": "LIQ-990001", "agrupaciones": "AGR-990001",
               "pagos": "PAG-990001", "ordenes_pago": "OP-990001", "perfiles": "", "acceso_cliente": "",
               "gastos": "GTO-EJ1", "mantenimientos": "MTO-EJ1"}

SI = {"SI", "SÍ", "S", "TRUE", "1", "X", "YES", "Y"}
NO = {"NO", "N", "FALSE", "0"}
RE_FECHA = re.compile(r"^\d{4}-\d{2}-\d{2}$")
RE_FECHA_HORA = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}$")
RE_UUID = re.compile(r"^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}$")

# Mensajes claros para los índices únicos de la base.
DUPLICADOS = [
    ("uq_viaje_duplicado", "ya hay un viaje con la misma fecha, piloto y ruta"),
    ("uq_tarifa_activa", "ya hay una tarifa activa para ese cliente, proyecto, ruta, capacidad y vigencia"),
    ("uq_ruta_codigo", "ya existe una ruta con ese código"),
    ("uq_factura_numero", "ya hay una factura activa con ese número"),
    ("clientes_cliente_key", "ya existe un cliente con esa razón social"),
    ("uq_clientes_nombre", "ya existe un cliente con ese nombre"),
    ("uq_clientes_nit", "ya existe un cliente fiscal con ese NIT"),
    ("perfiles_email_key", "ya existe un perfil con ese correo"),
    ("perfiles_usuario_key", "ya existe un perfil con ese usuario"),
]


def url_bd():
    if os.environ.get("DATABASE_URL"):
        return os.environ["DATABASE_URL"]
    txt = open(CONEXION, encoding="utf-8", errors="replace").read()
    m = re.search(r'(postgres(?:ql)?://[^\s"\'<>]+)', txt)
    if not m:
        raise SystemExit("No encontré DATABASE_URL en " + CONEXION)
    return m.group(1)


def introspect(cur):
    meta = {}
    for t in ORDEN:
        cur.execute("""select column_name, udt_name, is_nullable, column_default
                       from information_schema.columns where table_schema='master' and table_name=%s
                       order by ordinal_position""", (t,))
        cols = {r[0]: {"udt": r[1], "nul": r[2] == "YES", "dfl": r[3]} for r in cur.fetchall()}
        cur.execute("""select a.attname from pg_index i
                       join pg_attribute a on a.attrelid=i.indrelid and a.attnum=any(i.indkey)
                       where i.indrelid=('master.%s')::regclass and i.indisprimary""" % t)
        pk = [r[0] for r in cur.fetchall()]
        cur.execute("""select conname, pg_get_constraintdef(oid) from pg_constraint
                       where conrelid=('master.%s')::regclass and contype='c'""" % t)
        enums = {}
        for _, d in cur.fetchall():
            d = re.sub(r"^CHECK\s*", "", d)
            m = re.match(r"\(*(\w+) = ANY \(ARRAY\[(.*)\]\)\)", d)
            if m:
                vals = re.findall(r"'([^']*)'", m.group(2))
                if m.group(1) != "estado" or len(vals) <= 8:
                    enums[m.group(1)] = vals
        cur.execute("""select a.attname, confrelid::regclass::text, af.attname
                       from pg_constraint c
                       join pg_attribute a on a.attrelid=c.conrelid and a.attnum=c.conkey[1]
                       join pg_attribute af on af.attrelid=c.confrelid and af.attnum=c.confkey[1]
                       where c.conrelid=('master.%s')::regclass and c.contype='f' and cardinality(c.conkey)=1""" % t)
        fks = {}
        for col, ref, refcol in cur.fetchall():
            rt = ref.split(".")[-1]
            fks[col] = (rt, refcol)
        meta[t] = {"cols": cols, "pk": pk, "enums": enums, "fks": fks}
    return meta


_refcache = {}


def limpiar_cache():
    _refcache.clear()


def refset(cur, tabla, col):
    k = (tabla, col)
    if k not in _refcache:
        if tabla == "users":
            cur.execute("select lower(email) from auth.users")
        else:
            cur.execute('select "%s"::text from master."%s"' % (col, tabla))
        _refcache[k] = {r[0] for r in cur.fetchall()}
    return _refcache[k]


def obj(v):
    """Devuelve el valor Python de un Jsonb de psycopg (o el valor tal cual)."""
    return getattr(v, "obj", v)


def fecha_texto(v):
    """Excel entrega fechas como datetime; la base las guarda como texto ISO."""
    if isinstance(v, datetime):
        if v.hour or v.minute or v.second:
            return v.strftime("%Y-%m-%dT%H:%M")
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    return v


def convertir(val, col, meta_col):
    """Devuelve (valor_python, error)."""
    if val is None:
        return None, None
    udt = meta_col["udt"]
    val = fecha_texto(val)
    if isinstance(val, str):
        val = val.strip()
        if val == "":
            return None, None
    try:
        if udt == "_text":
            if isinstance(val, list):
                return [str(x).strip() for x in val], None
            return [x.strip() for x in re.split(r"[|;]", str(val)) if x.strip()], None
        if udt == "jsonb":
            v = val
            if isinstance(v, str):
                v = json.loads(v)
            if isinstance(v, str):
                v = json.loads(v)          # corrige el doble codificado antiguo
            return Jsonb(v), None
        if udt == "bool":
            s = str(val).strip().upper()
            if s in SI:
                return True, None
            if s in NO:
                return False, None
            return None, "booleano inválido: %r (usa SI/NO)" % val
        if udt in ("int4", "int8"):
            return int(Decimal(str(val).replace(",", ""))), None
        if udt == "numeric":
            return Decimal(str(val).replace(",", "")), None
        if udt == "timestamptz":
            s = str(val).strip().replace(" ", "T")
            if not re.match(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}(:\d{2})?$", s):
                return None, "fecha y hora inválida: %r (usa AAAA-MM-DDTHH:MM)" % val
            return s, None
        if udt == "uuid":
            return str(val).strip(), None
        if isinstance(val, float) and val.is_integer():
            return str(int(val)), None
        return str(val), None
    except (ValueError, InvalidOperation, json.JSONDecodeError) as e:
        return None, "valor inválido (%s): %r" % (e.__class__.__name__, val)


def amigable(msg):
    bajo = msg.lower()
    for indice, texto in DUPLICADOS:
        if indice in bajo:
            return texto[0].upper() + texto[1:]
    if "duplicate key" in bajo:
        return "ya existe una fila con esa misma clave"
    if "violates foreign key" in bajo:
        return "una referencia no existe (revisa la hoja de la tabla relacionada)"
    if "violates check constraint" in bajo:
        m = re.search(r'constraint "([^"]+)"', msg)
        return "no cumple la regla de la base" + (" (%s)" % m.group(1) if m else "")
    return msg


def clave_de(t, fila, meta):
    pk = meta[t]["pk"]
    if not pk or any(p not in fila for p in pk):
        return None
    return "|".join(str(fila[p]) for p in pk)


def dato(cur, sql, args=()):
    cur.execute(sql, args)
    r = cur.fetchone()
    return r[0] if r else None


def fila_dato(cur, sql, args=()):
    cur.execute(sql, args)
    return cur.fetchone()


def validar(t, fila, meta, cur, verbose):
    """fila: {col: valor} · devuelve (fila_lista, errores[])."""
    errs = []
    m = meta[t]
    if t == "perfiles":
        if not fila.get("id") and fila.get("email"):
            cur.execute("select id::text from auth.users where lower(email)=lower(%s)", (fila["email"],))
            r = cur.fetchone()
            if r:
                fila["id"] = r[0]
            else:
                errs.append("no existe el login %s en Supabase Auth: créalo primero" % fila["email"])
        if fila.get("id") and not RE_UUID.match(str(fila["id"])):
            errs.append("id no es un UUID")
        if fila.get("id") and RE_UUID.match(str(fila["id"])):
            cur.execute("select 1 from auth.users where id=%s", (fila["id"],))
            if not cur.fetchone():
                errs.append("el UUID %s no existe en Supabase Auth: crea primero el login" % fila["id"])
    if t == "acceso_cliente":
        u = fila.get("user_id")
        if u and not RE_UUID.match(str(u)):
            cur.execute("select id::text from master.perfiles where lower(email)=lower(%s) or usuario=%s", (u, u))
            r = cur.fetchone()
            if r:
                fila["user_id"] = r[0]
            else:
                errs.append("user_id %r no es UUID ni un correo/usuario existente" % u)
    # obligatorias (NOT NULL sin default)
    for col, cm in m["cols"].items():
        if col not in fila and (col in m["pk"] or (not cm["nul"] and not cm["dfl"])):
            errs.append("falta la columna obligatoria %s" % col)
    # enums y formatos
    for col, v in list(fila.items()):
        if col not in m["cols"]:
            errs.append("columna desconocida: %s" % col)
            continue
        if col in m["enums"] and v is not None and str(v) not in m["enums"][col]:
            errs.append("%s=%r no es válido (usa: %s)" % (col, v, " · ".join(x for x in m["enums"][col] if x)))
        if t == "viajes" and col == "fecha_plan" and v is not None and not RE_FECHA_HORA.match(str(v)):
            errs.append("fecha_plan debe ser AAAA-MM-DDTHH:MM")
        if t == "viajes" and col.startswith("t_") and v is not None and not RE_FECHA_HORA.match(str(v)):
            errs.append("%s debe ser AAAA-MM-DDTHH:MM o vacío" % col)
        if t in ("tarifas",) and col in ("vigente_desde", "vigente_hasta") and v not in (None, "") and not RE_FECHA.match(str(v)):
            errs.append("%s debe ser AAAA-MM-DD" % col)
        if t == "mem" and col == "fecha" and v is not None and not RE_FECHA.match(str(v)):
            errs.append("mem.fecha debe ser AAAA-MM-DD")
        if t == "pagos" and col == "fecha" and v is not None and not RE_FECHA.match(str(v)):
            errs.append("pagos.fecha debe ser AAAA-MM-DD")
        if col in ("importe", "monto", "precio_diesel") and v is not None and (v <= 0):
            errs.append("%s debe ser mayor que cero" % col)
    # llaves foráneas
    for col, (rt, rc) in m["fks"].items():
        v = fila.get(col)
        if v in (None, ""):
            continue
        if rt == "users":
            continue
        if str(v) not in refset(cur, rt, rc):
            errs.append("%s=%r no existe en %s.%s" % (col, v, rt, rc))
    # reglas de negocio por tabla
    if t == "rutas" and fila.get("estado") in ("ACTIVA", "EN_TARIFACION", "PENDIENTE_APROBACION"):
        for col in ("codigo", "origen", "destino"):
            if not fila.get(col):
                errs.append("una ruta %s exige %s" % (fila["estado"], col))
    if t == "rutas" and fila.get("estado") != "RECHAZADA" and fila.get("codigo"):
        otro = dato(cur, "select id from master.rutas where codigo=%s and id<>%s limit 1", (fila["codigo"], fila.get("id", "")))
        if otro:
            errs.append("el código %r ya lo usa la ruta %s" % (fila["codigo"], otro))
    if t == "clientes":
        if fila.get("id") and not re.match(r"^[A-Z0-9_-]{2,20}$", str(fila["id"])):
            errs.append("id de cliente inválido (2-20: A-Z, 0-9, _ o -)")
        if not fila.get("proyectos"):
            errs.append("proyectos no puede quedar vacío: sin proyectos no se pueden crear viajes ni tarifas")
        if fila.get("cliente"):
            otro = dato(cur, "select id from master.clientes where lower(cliente)=lower(%s) and id<>%s limit 1", (fila["cliente"], fila.get("id", "")))
            if otro:
                errs.append("ya existe el cliente %r con id %s" % (fila["cliente"], otro))
    if t == "flota":
        ex = obj(fila.get("extra"))
        if isinstance(ex, dict):
            prop = str(ex.get("propiedad") or "").upper()
            if prop and prop not in ("INTERNA", "OPL"):
                errs.append("flota.extra.propiedad debe ser INTERNA u OPL")
            cli = ex.get("clienteId")
            if cli and str(cli) not in refset(cur, "clientes", "id"):
                errs.append("flota.extra.clienteId=%r no existe en clientes" % cli)
            opl = ex.get("opl")
            if opl:
                cur.execute("select valores from master.catalogos where clave='opl'")
                r = cur.fetchone()
                vals = r[0] if r else []
                if isinstance(vals, str):
                    try: vals = json.loads(vals)
                    except Exception: vals = []
                if opl not in (vals or []):
                    errs.append("flota.extra.opl=%r no está en la lista opl (%s)" % (opl, ", ".join(vals or [])))
    if t == "clientes_fiscal" and fila.get("nit"):
        otro = dato(cur, "select cliente_id from master.clientes_fiscal where nit=%s and cliente_id<>%s limit 1", (fila["nit"], fila.get("cliente_id", "")))
        if otro:
            errs.append("el NIT %r ya está en el cliente %s" % (fila["nit"], otro))
    if t == "tarifas":
        if fila.get("activa", True) and fila.get("vigente_desde"):
            otro = dato(cur, """select id from master.tarifas
                                where cliente_id=%s and proyecto=%s and ruta_id=%s and capacidad=%s
                                  and vigente_desde=%s and activa and id<>%s limit 1""",
                        (fila.get("cliente_id"), fila.get("proyecto"), fila.get("ruta_id"), fila.get("capacidad"),
                         fila.get("vigente_desde"), fila.get("id", "")))
            if otro:
                errs.append("ya hay una tarifa activa igual (%s); desactiva la anterior o cambia la vigencia" % otro)
        if fila.get("ruta_id"):
            r = fila_dato(cur, "select cliente_id, proyecto from master.rutas where id=%s", (fila["ruta_id"],))
            if r and (r[0] != fila.get("cliente_id") or r[1] != fila.get("proyecto")):
                errs.append("la ruta %s es de %s / %s, no de %s / %s" % (fila["ruta_id"], r[0], r[1], fila.get("cliente_id"), fila.get("proyecto")))
    if t == "viajes":
        if fila.get("id") and not re.match(r"^\d{6,}$", str(fila["id"])):
            errs.append("id de viaje inválido: debe tener 6 dígitos o más (ej. 102656)")
        if fila.get("fecha_plan") and fila.get("piloto_id") and fila.get("ruta_id"):
            otro = dato(cur, """select id from master.viajes
                                where fecha_plan=%s and piloto_id=%s and ruta_id=%s and id<>%s limit 1""",
                        (fila["fecha_plan"], fila["piloto_id"], fila["ruta_id"], fila.get("id", "")))
            if otro:
                errs.append("viaje duplicado: ya existe %s con la misma fecha, piloto y ruta" % otro)
        proys = None
        if fila.get("cliente_id"):
            proys = dato(cur, "select proyectos from master.clientes where id=%s", (fila["cliente_id"],))
        if proys and fila.get("proyecto_cat") and fila["proyecto_cat"] not in proys:
            errs.append("proyecto_cat=%r no está en los proyectos del cliente (%s)" % (fila["proyecto_cat"], ", ".join(proys)))
        if fila.get("ruta_id"):
            r = fila_dato(cur, "select cliente_id, proyecto from master.rutas where id=%s", (fila["ruta_id"],))
            if r and (r[0] != fila.get("cliente_id") or r[1] != fila.get("proyecto_cat")):
                errs.append("la ruta %s es de %s / %s, no de %s / %s" % (fila["ruta_id"], r[0], r[1], fila.get("cliente_id"), fila.get("proyecto_cat")))
        if verbose and fila.get("piloto_id") and fila.get("cliente_id"):
            pc = dato(cur, "select cliente_id from master.pilotos where id=%s", (fila["piloto_id"],))
            if pc and pc != fila["cliente_id"]:
                print("      aviso: el piloto %s es del cliente %s (el viaje es de %s)" % (fila["piloto_id"], pc, fila["cliente_id"]))
    if t == "agrupaciones":
        if fila.get("fase") == "EMITIDA" and (not obj(fila.get("emision")) or fila.get("importe_facturado") in (None, "")):
            errs.append("fase EMITIDA exige emision e importe_facturado")
        if fila.get("fase") not in (None, "PREPARADA") and not fila.get("numero_factura"):
            errs.append("fase %s exige numero_factura" % fila.get("fase"))
        if fila.get("fase") in ("ENVIADA", "EMITIDA") and not obj(fila.get("envio")):
            errs.append("fase %s exige envio (quién y cuándo se envió)" % fila.get("fase"))
        if fila.get("numero_factura") and fila.get("estado", "ACTIVA") == "ACTIVA":
            otro = dato(cur, "select id from master.agrupaciones where lower(numero_factura)=lower(%s) and estado='ACTIVA' and id<>%s limit 1",
                        (fila["numero_factura"], fila.get("id", "")))
            if otro:
                errs.append("el número de factura %r ya está en la factura activa %s" % (fila["numero_factura"], otro))
        ids = fila.get("viaje_ids") or []
        if ids:
            cur.execute("select id from master.viajes where id = any(%s)", (ids,))
            hay = {r[0] for r in cur.fetchall()}
            faltan = [x for x in ids if x not in hay]
            if faltan:
                errs.append("viaje_ids no existen: " + ", ".join(faltan[:8]) + ("…" if len(faltan) > 8 else ""))
            if fila.get("cliente_id"):
                cur.execute("select count(*) from master.viajes where id = any(%s) and cliente_id<>%s", (ids, fila["cliente_id"]))
                if cur.fetchone()[0]:
                    errs.append("hay viajes de otro cliente en esta factura")
            imp = obj(fila.get("importes_viaje")) or {}
            if isinstance(imp, dict):
                fuera = [k for k in imp if k not in set(ids)]
                if fuera:
                    errs.append("importes_viaje tiene viajes que no están en viaje_ids: " + ", ".join(fuera[:6]))
    if t == "ordenes_pago":
        if fila.get("estado", "ACTIVA") == "ACTIVA" and not obj(fila.get("emision")):
            errs.append("una orden ACTIVA exige emision")
        if fila.get("fase") == "PAGADA" and not obj(fila.get("transferencia")):
            errs.append("fase PAGADA exige transferencia")
        ids = fila.get("viaje_ids") or []
        if ids:
            cur.execute("select id from master.viajes where id = any(%s)", (ids,))
            hay = {r[0] for r in cur.fetchall()}
            faltan = [x for x in ids if x not in hay]
            if faltan:
                errs.append("viaje_ids no existen: " + ", ".join(faltan[:8]) + ("…" if len(faltan) > 8 else ""))
    if t == "liquidaciones":
        if fila.get("decision") == "DEVUELTA" and len(str(fila.get("comentario") or "")) < 5:
            errs.append("DEVUELTA exige comentario de al menos 5 caracteres")
        if fila.get("ciclo") is not None and int(fila["ciclo"]) < 1:
            errs.append("ciclo debe ser 1 o mayor")
    return fila, errs


def cargar(cur, ws, t, meta, aplicar, probar, verbose):
    m = meta[t]
    heads = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    # un encabezado vacío en medio desplazaría las columnas: mejor detectarlo
    ultimo = max([i for i, h in enumerate(heads, 1) if h] or [0])
    huecos = [i for i, h in enumerate(heads[:ultimo], 1) if not h]
    if huecos:
        return 0, ["hay columnas sin encabezado en la fila 1 (posición %s)" % ", ".join(map(str, huecos))]
    heads = heads[:ultimo]
    desconocidas = [h for h in heads if h not in m["cols"]]
    if desconocidas:
        return 0, ["columnas que no existen en la tabla: " + ", ".join(desconocidas)]
    filas, errores = [], []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, i + 1).value for i in range(len(heads))]
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in vals):
            continue
        if probar:
            if r != 2:
                continue                     # solo la fila de ejemplo
        elif r == 2:
            continue                         # la fila 2 es la de ejemplo: los datos van desde la 3
        fila = {}
        mal = False
        for h, v in zip(heads, vals):
            pv, err = convertir(v, h, m["cols"][h])
            if err:
                errores.append("fila %d · %s: %s" % (r, h, err))
                mal = True
            elif pv is not None:
                fila[h] = pv
        if mal:
            continue
        if probar and heads and t in PROBE_FIRST:
            fila[heads[0]] = PROBE_FIRST[t]
        filas.append((r, fila))
    insertadas = 0
    vistas = {}
    for r, fila in filas:
        fila, errs = validar(t, fila, meta, cur, verbose)
        if errs:
            errores.append("fila %d: %s" % (r, " · ".join(errs)))
            continue
        k = clave_de(t, fila, meta)
        if k:
            if k in vistas:
                errores.append("fila %d: clave repetida en la hoja (igual que la fila %d): %s" % (r, vistas[k], k))
                continue
            vistas[k] = r
        cols = list(fila.keys())
        if not cols:
            errores.append("fila %d: no tiene ningún dato" % r)
            continue
        sql = 'insert into master."%s" (%s) values (%s)' % (
            t, ", ".join('"%s"' % c for c in cols), ", ".join(["%s"] * len(cols)))
        cur.execute("savepoint fila")
        try:
            cur.execute(sql, [fila[c] for c in cols])
            cur.execute("release savepoint fila")
            insertadas += 1
        except Exception as e:
            cur.execute("rollback to savepoint fila")
            msg = str(e).strip().splitlines()[0]
            if probar and "duplicate key" in msg:
                continue
            errores.append("fila %d: %s" % (r, amigable(msg)[:300]))
    return insertadas, errores


def cargar_listas(cur, ws, tipo, probar, reemplazar):
    """Hoja catalogos → master.catalogos (listas del sistema).
    Se guardan como la app: un arreglo JSON dentro del campo `valores`."""
    cur.execute("select clave, valores from master.catalogos")
    actuales = {}
    for k, v in cur.fetchall():
        if isinstance(v, str):
            try:
                v = json.loads(v)
            except Exception:
                v = []
        actuales[k] = [str(x) for x in (v or [])]
    heads = [str(c.value).strip() for c in ws[1] if c.value]
    filas = []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, i + 1).value for i in range(len(heads))]
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in vals):
            continue
        filas.append((r, vals))
    if probar and not filas:
        filas = [(2, ["opl", "EJEMPLO-OPL"])]
    errores, grupos = [], {}
    for r, vals in filas:
        lista = str(vals[0] or "").strip() if len(vals) > 0 else ""
        valor = str(vals[1] or "").strip() if len(vals) > 1 else ""
        if not lista:
            errores.append("fila %d: falta la lista" % r)
            continue
        if lista not in CLAVES_LISTA:
            errores.append("fila %d: lista %r desconocida (usa %s)" % (r, lista, " · ".join(sorted(CLAVES_LISTA))))
            continue
        if not valor:
            errores.append("fila %d: falta el valor" % r)
            continue
        grupos.setdefault(lista, []).append(valor)
    nuevos_total = 0
    for lista, vals in grupos.items():
        unicos = list(dict.fromkeys(vals))
        base = [] if reemplazar else list(actuales.get(lista, []))
        nuevos = [v for v in unicos if v not in base]
        nueva = base + nuevos
        if nueva != actuales.get(lista, []):
            cur.execute("""insert into master.catalogos (clave, etiqueta, valores, actualizado_por, actualizado_en)
                           values (%s, %s, %s::jsonb, 'plantilla', now())
                           on conflict (clave) do update set valores = excluded.valores,
                             actualizado_por = 'plantilla', actualizado_en = now()""",
                        (lista, lista, json.dumps(nueva, ensure_ascii=False)))
        nuevos_total += len(nuevos)
        print("      %s: %d valores%s" % (lista, len(nueva), " (%d nuevos)" % len(nuevos) if nuevos else ""))
    return nuevos_total, errores


def borrar_datos(cur, aplicar):
    """Deja el esquema listo para los datos reales. Borra las tablas de negocio en orden.
    Solo se ejecuta dentro de la transacción de carga: si algo falla, todo vuelve."""
    total = 0
    for t in PURGA:
        cur.execute('select count(*) from master."%s"' % t)
        n = cur.fetchone()[0]
        total += n
        if n:
            print("      %-18s %6d filas" % (t, n))
    print("      total a borrar: %d filas (los usuarios y las listas se conservan)" % total)
    marca = datetime.now().strftime("%Y%m%d_%H%M%S")
    carpeta = os.path.join(RESPALDOS, "antes_de_carga_" + marca)
    if aplicar:
        os.makedirs(carpeta, exist_ok=True)
        for t in PURGA:
            cur.execute('select count(*) from master."%s"' % t)
            if not cur.fetchone()[0]:
                continue
            with open(os.path.join(carpeta, t + ".jsonl"), "wb") as f:
                with cur.copy('copy (select row_to_json(x)::text from master."%s" x) to stdout' % t) as cp:
                    for chunk in cp:
                        f.write(chunk)
                        f.write(bytes([10]))
        print("      respaldo: " + carpeta)
    for t in PURGA:
        cur.execute('alter table master."%s" disable trigger user' % t)
    try:
        for t in PURGA:
            cur.execute('delete from master."%s"' % t)
    finally:
        for t in PURGA:
            cur.execute('alter table master."%s" enable trigger user' % t)
    limpiar_cache()
    return total


def sincronizar_secuencias(cur):
    cur.execute("select clave, valor from master.secuencias")
    actuales = dict(cur.fetchall())
    calc = {}
    pares = [("trip", "viajes", "id", r"^\d+$"), ("tar", "tarifas", "id", r"^TAR-\d+$"),
             ("rut", "rutas", "id", r"^R-\d+$"), ("agr", "agrupaciones", "id", r"^AGR-\d+$"),
             ("op", "ordenes_pago", "id", r"^OP-\d+$"), ("pag", "pagos", "id", r"^PAG-\d+$"),
             ("liq", "liquidaciones", "id", r"^LIQ-\d+$"), ("evi", "evidencias", "id", r"^EVI-\d+$"),
             ("pil", "pilotos", "id", r"^P-\d+$")]
    for clave, tabla, col, patron in pares:
        cur.execute("""select coalesce(max(nullif(regexp_replace(%s, '\\D', '', 'g'), '')::bigint), 0)
                       from master.%s where %s ~ %%s""" % (col, tabla, col), (patron,))
        calc[clave] = int(cur.fetchone()[0]) + 1
    cur.execute("""select coalesce(max(nullif(regexp_replace(c->>'id', '\\D', '', 'g'), '')::bigint), 0)
                   from master.agrupaciones a, jsonb_array_elements(a.cobros) c""")
    calc["cob"] = int(cur.fetchone()[0]) + 1
    cambios = []
    for clave, valor in calc.items():
        if valor > 1 and valor > int(actuales.get(clave, 0)):
            cur.execute("""insert into master.secuencias (clave, valor) values (%s, %s)
                           on conflict (clave) do update set valor = excluded.valor""", (clave, valor))
            cambios.append("%s=%d" % (clave, valor))
    return cambios


def main():
    ap = argparse.ArgumentParser(description="Carga Plantillas_Master.xlsx en master.*")
    ap.add_argument("--archivo", default=os.path.join(BASE, "Plantillas_Master.xlsx"))
    ap.add_argument("--aplicar", action="store_true", help="confirma los cambios (si no, revierte)")
    ap.add_argument("--simular", action="store_true", help="inserta y revierte (por defecto)")
    ap.add_argument("--probar", action="store_true", help="prueba con la fila de ejemplo y revierte")
    ap.add_argument("--tabla", default=None, help="carga solo esta hoja")
    ap.add_argument("--reemplazar", action="store_true",
                    help="borra primero los datos de negocio (clientes, flota, viajes, facturas...) y deja solo usuarios y listas")
    ap.add_argument("--si-reemplazar", action="store_true", help="confirma el borrado cuando se usa --aplicar")
    ap.add_argument("--reemplazar-listas", action="store_true",
                    help="la hoja catalogos sustituye las listas actuales; por defecto se suman valores")
    ap.add_argument("--verbose", action="store_true")
    a = ap.parse_args()
    aplicar = a.aplicar and not a.probar
    probar = a.probar
    if a.reemplazar and a.tabla:
        raise SystemExit("--reemplazar no se combina con --tabla: borra todo y debe cargar el libro completo.")
    if a.reemplazar and aplicar and not a.si_reemplazar:
        raise SystemExit("Para borrar los datos actuales y cargar los reales agrega también --si-reemplazar.")
    wb = load_workbook(a.archivo, data_only=True, read_only=False)
    if a.tabla and a.tabla not in SECUENCIA:
        raise SystemExit("Tabla desconocida. Opciones: " + ", ".join(SECUENCIA))
    with psycopg.connect(url_bd()) as cn:
        cur = cn.cursor()
        meta = introspect(cur)
        total_ok, total_err, resumen = 0, 0, []
        try:
            if a.reemplazar:
                print("· reemplazo de datos de negocio")
                borrar_datos(cur, aplicar)
            for t in ([a.tabla] if a.tabla else SECUENCIA):
                if probar and t in SALTAR_EN_PRUEBA:
                    continue
                if t not in wb.sheetnames:
                    print("· %-18s (sin hoja, se omite)" % t)
                    continue
                if t in LISTAS:
                    n, errs = cargar_listas(cur, wb[t], t, probar, a.reemplazar_listas)
                    total_ok += n
                    total_err += len(errs)
                    resumen.append((t, n, len(errs)))
                    print("· %-18s %4d valores · %d error(es)" % (t, n, len(errs)))
                    for e in errs[:8]:
                        print("      " + e)
                    continue
                ok, errs = cargar(cur, wb[t], t, meta, aplicar, probar, a.verbose)
                total_ok += ok
                total_err += len(errs)
                resumen.append((t, ok, len(errs)))
                print("· %-18s %4d filas · %d error(es)" % (t, ok, len(errs)))
                for e in errs[:8]:
                    print("      " + e)
                if len(errs) > 8:
                    print("      … y %d más" % (len(errs) - 8))
            if not probar:
                try:
                    cur.execute("select nombre from master.proveedores where activo")
                    provs = {r[0] for r in cur.fetchall()}
                    cur.execute("select valores from master.catalogos where clave='opl'")
                    r = cur.fetchone()
                    opl = json.loads(r[0]) if r and isinstance(r[0], str) else (r[0] if r else [])
                    faltan = sorted(provs - set(opl or []))
                    if faltan:
                        print("  aviso: proveedores que no están en la lista opl: " + ", ".join(faltan))
                except Exception:
                    pass
            if total_err:
                cn.rollback()
                print("\n%d errores: no se guardó nada." % total_err)
                raise SystemExit(1)
            if aplicar:
                cambios = sincronizar_secuencias(cur)
                cn.commit()
                print("\nCARGADO: %d filas. Contadores actualizados: %s" %
                      (total_ok, ", ".join(cambios) if cambios else "sin cambios"))
            else:
                cn.rollback()
                print("\nSIMULACIÓN OK: %d filas válidas; se revirtió todo.%s" %
                      (total_ok, " (incluye la fila de EJEMPLO de cada hoja)" if probar else ""))
        except SystemExit:
            raise
        except Exception as e:
            cn.rollback()
            print("\nERROR inesperado, se revirtió todo:", e)
            raise


if __name__ == "__main__":
    main()
